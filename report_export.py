"""User-triggered Word and PDF export for the latest assistant response."""
from datetime import datetime, timezone
from io import BytesIO
import html
import os
from pathlib import Path

MAX_REPORT_CHARS = 60000


def _bounded_text(text):
    if not isinstance(text, str) or not text.strip():
        raise ValueError("Rapor metni boş olamaz.")
    if len(text) > MAX_REPORT_CHARS:
        raise ValueError("Yanıt 60.000 karakter rapor sınırını aşıyor.")
    return text


def make_docx(text, title="Nesil AI Raporu"):
    from docx import Document

    text = _bounded_text(text)
    document = Document()
    document.add_heading(title, level=0)
    document.add_paragraph("Oluşturulma zamanı (UTC): " + datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"))
    document.add_paragraph("Nesil AI yanıtı — önemli kararlar öncesinde içeriği ve kaynakları doğrulayın.")
    for line in text.splitlines():
        document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _unicode_font():
    candidates = [
        os.environ.get("NESIL_PDF_FONT", ""),
        r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\calibri.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return candidate
    raise ValueError("Türkçe PDF için Unicode TTF yazı tipi bulunamadı. NESIL_PDF_FONT ayarlayın.")


def make_pdf(text, title="Nesil AI Raporu"):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    text = _bounded_text(text)
    font_path = _unicode_font()
    font_name = "NesilUnicode"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, font_path))
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("NesilTitle", parent=styles["Title"], fontName=font_name,
                                 textColor=colors.HexColor("#12346b"), alignment=TA_LEFT, leading=25)
    body_style = ParagraphStyle("NesilBody", parent=styles["BodyText"], fontName=font_name,
                                fontSize=10, leading=15, spaceAfter=6, wordWrap="LTR")
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, rightMargin=48, leftMargin=48, topMargin=48, bottomMargin=48)
    story = [Paragraph(html.escape(title), title_style), Spacer(1, 8),
             Paragraph("Oluşturulma zamanı (UTC): " + datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"), body_style),
             Paragraph("Nesil AI yanıtı — önemli kararlar öncesinde içeriği ve kaynakları doğrulayın.", body_style),
             Spacer(1, 8)]
    for line in text.splitlines():
        story.append(Paragraph(html.escape(line) if line else "&nbsp;", body_style))
    document.build(story)
    return output.getvalue()
