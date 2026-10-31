"""Bounded document extraction with explicit scope and per-file status."""
import csv
import io
import zipfile
import re
from dataclasses import dataclass

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TOTAL_BYTES = 20 * 1024 * 1024
MAX_FILES = 5
MAX_TEXT = 240000
ALLOWED = ("txt", "md", "json", "py", "log", "csv", "pdf", "docx", "xlsx", "xlsm", "pptx",
           "png", "jpg", "jpeg", "tif", "tiff", "bmp")
IMAGE_TYPES = {"png", "jpg", "jpeg", "tif", "tiff", "bmp"}


@dataclass
class FileResult:
    name: str
    text: str = ""
    status: str = "okunamadı"
    scope: str = ""
    chart_data: dict | None = None


def extract_file(upload, ocr=False):
    name = upload.name
    if upload.size > MAX_FILE_BYTES:
        return FileResult(name, scope="10 MB sınırı aşıldı.")
    ext = name.lower().rsplit(".", 1)[-1]
    if ext not in ALLOWED:
        return FileResult(name, scope="Desteklenmeyen dosya türü.")
    raw = upload.getvalue()
    if len(raw) > MAX_FILE_BYTES:
        return FileResult(name, scope="10 MB sınırı aşıldı.")
    parts, scope = [], ""
    sampled = False
    chart_headers, chart_rows, chart_label = [], [], ""
    try:
        if ext in ("docx", "xlsx", "xlsm", "pptx"):
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                entries = archive.infolist()
                if len(entries) > 2000 or sum(e.file_size for e in entries) > 50 * 1024 * 1024:
                    return FileResult(name, scope="Belgenin açılmış boyutu veya öğe sayısı güvenlik sınırını aşıyor.")
        if ext in ("txt", "md", "json", "py", "log"):
            text = raw.decode("utf-8-sig", errors="replace")
            parts = [text[:MAX_TEXT + 1]]
            scope = "UTF-8 metni; çözülemeyen karakterler değiştirilir."
        elif ext == "csv":
            decoded = raw.decode("utf-8-sig", errors="replace")
            try:
                dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(io.StringIO(decoded), dialect)
            total = 0
            for index, row in enumerate(reader, 1):
                if index > 5000:
                    sampled = True
                    break
                if len(row) > 50:
                    sampled = True
                if index == 1:
                    chart_headers = row[:20]
                    chart_label = "CSV"
                elif index <= 501:
                    chart_rows.append(row[:20])
                line = f"Satır {index}: " + " | ".join(cell[:500] for cell in row[:50])
                sampled = sampled or any(len(cell) > 500 for cell in row[:50])
                parts.append(line)
                total += len(line)
                if total > MAX_TEXT:
                    sampled = True
                    break
            scope = "En fazla 5.000 satır, 50 sütun ve hücre başına 500 karakter."
        elif ext == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw))
            sampled = len(reader.pages) > 200
            total = 0
            ocr_count, ocr_failed, empty_pages = 0, False, 0
            for index, page in enumerate(reader.pages[:200], 1):
                page_text = page.extract_text() or ""
                native_text = bool(page_text.strip())
                if not page_text.strip():
                    empty_pages += 1
                    if ocr and ocr_count < 20 and not ocr_failed:
                        try:
                            from ocr_reader import ocr_page
                            page_text = ocr_page(raw, index - 1)
                            ocr_count += 1
                        except Exception:
                            ocr_failed = True
                text = f"[Sayfa {index}{' OCR' if ocr and page_text.strip() and not native_text else ''}]\n" + page_text
                parts.append(text[:MAX_TEXT + 1])
                total += len(text)
                if total > MAX_TEXT:
                    sampled = True
                    break
            scope = f"PDF: toplam {len(reader.pages)} sayfa; en fazla ilk 200 sayfa ve 240.000 karakter. OCR yok."
            if ocr:
                scope = scope.replace("OCR yok.", f"Yerel OCR: {ocr_count} sayfa; en fazla 20 taranmış sayfa. OCR metni hatalı olabilir.")
                if ocr_failed:
                    scope += " OCR kullanılamadı: isteğe bağlı paketleri, Tesseract ve dil paketlerini kontrol edin."
            sampled = sampled or empty_pages > ocr_count or ocr_failed
        elif ext == "docx":
            from docx import Document
            doc = Document(io.BytesIO(raw))
            total = 0
            for paragraph_index, paragraph in enumerate(doc.paragraphs, 1):
                if paragraph.text.strip():
                    parts.append(f"[Paragraf {paragraph_index}]\n" + paragraph.text[:MAX_TEXT + 1])
                    total += len(paragraph.text)
                if total > MAX_TEXT:
                    break
            for index, table in enumerate(doc.tables[:30], 1):
                if total > MAX_TEXT:
                    break
                for row_index, row in enumerate(table.rows[:1000], 1):
                    text = f"[Tablo {index}, satır {row_index}] " + " | ".join(c.text[:500] for c in row.cells[:50])
                    parts.append(text)
                    total += len(text)
                    if total > MAX_TEXT:
                        break
            sampled = True
            scope = "Word ana metni ve en fazla 30 tablodan ilk 1.000 satır/50 sütun; toplam 240.000 karakter. Üstbilgi, metin kutusu ve görseller kapsam dışı."
        elif ext == "pptx":
            from pptx import Presentation
            presentation = Presentation(io.BytesIO(raw))
            sampled = len(presentation.slides) > 200
            total = 0
            for slide_index in range(min(len(presentation.slides), 200)):
                slide_number = slide_index + 1
                slide = presentation.slides[slide_index]
                slide_parts = []
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        value = shape.text.strip()
                        if value:
                            slide_parts.append(value[:5000])
                    if getattr(shape, "has_table", False):
                        for row_index in range(min(len(shape.table.rows), 200)):
                            row_number = row_index + 1
                            row = shape.table.rows[row_index]
                            cells = [row.cells[index].text.strip()[:500]
                                     for index in range(min(len(row.cells), 30))]
                            if any(cells):
                                slide_parts.append(f"Tablo satırı {row_number}: " + " | ".join(cells))
                slide_text = f"[Slayt {slide_number}]\n" + "\n".join(slide_parts)
                total += len(slide_text)
                parts.append(slide_text[:MAX_TEXT + 1])
                if total > MAX_TEXT:
                    sampled = True
                    break
            scope = "PowerPoint: ilk 200 slayt; slayt metinleri ve tablolar. Gömülü görsel, grafik serisi, animasyon ve konuşmacı notları kapsam dışı."
        elif ext in IMAGE_TYPES:
            if not ocr:
                return FileResult(name, scope="Görsel OCR için 'yerel OCR kullan' seçeneğini etkinleştirin.")
            from ocr_reader import ocr_image
            try:
                image_text = ocr_image(raw)
            except Exception as exc:
                return FileResult(name, scope=f"Görsel OCR kullanılamadı ({type(exc).__name__}). Tesseract ve Türkçe/İngilizce dil paketlerini kontrol edin.")
            if not image_text.strip():
                return FileResult(name, scope="Görselde okunabilir metin bulunamadı. Daha net bir görsel deneyin.")
            parts = ["[Görsel OCR]\n" + image_text]
            sampled = len(image_text) > MAX_TEXT
            scope = "Yerel OCR görselin ilk karesini okur; metin tanıma hatalı olabilir."
        else:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            try:
                total = 0
                for sheet_index, ws in enumerate(wb.worksheets[:20]):
                    if total >= MAX_TEXT:
                        break
                    parts.append(f"[Sayfa: {ws.title}]")
                    sheet_chars = 0
                    for row_index, row in enumerate(ws.iter_rows(min_row=1, max_row=2000, max_col=50, values_only=True), 1):
                        if sheet_index == 0 and row_index == 1:
                            chart_headers = list(row[:20])
                            chart_label = ws.title
                        elif sheet_index == 0 and row_index <= 501 and any(value is not None for value in row[:20]):
                            chart_rows.append(list(row[:20]))
                        values = [str(v)[:500] if v is not None else "" for v in row]
                        if any(values):
                            from openpyxl.utils import get_column_letter
                            line = f"[Sayfa: {ws.title}, satır {row_index}] " + " | ".join(f"{get_column_letter(i)}{row_index}={value}" for i, value in enumerate(values, 1) if value)
                            parts.append(line)
                            sheet_chars += len(line)
                            total += len(line)
                            if sheet_chars >= 60000 or total >= MAX_TEXT:
                                break
            finally:
                wb.close()
            sampled = True
            scope = "Excel: ilk 20 sayfadan en fazla 2.000 satır/50 sütun, sayfa başına yaklaşık 60.000 ve toplam 240.000 karakter. Formüller hesaplanmaz; kayıtlı sonuçlar boş/eski olabilir. Makrolar çalıştırılmaz."
        text = "\n".join(parts)
        sampled = sampled or len(text) > MAX_TEXT
        text = text[:MAX_TEXT]
        if not text.strip() or (ext == "pdf" and not any((p.split("\n", 1)[-1]).strip() for p in parts)):
            return FileResult(name, scope="Okunabilir metin bulunamadı. Taranmış PDF için OCR gerekebilir. " + scope)
        chart_data = None
        if chart_headers and chart_rows:
            headers = []
            used = set()
            for index, value in enumerate(chart_headers, 1):
                base = str(value).strip() if value is not None else ""
                base = base or f"Sütun {index}"
                label, suffix = base, 2
                while label in used:
                    label, suffix = f"{base} ({suffix})", suffix + 1
                used.add(label)
                headers.append(label)
            chart_rows = [row[:len(headers)] + [None] * max(0, len(headers) - len(row)) for row in chart_rows]
            chart_data = {"label": chart_label, "columns": headers, "rows": chart_rows}
        return FileResult(name, text, "kısmen okundu" if sampled else "okundu",
                          scope + f" Çıkarılan kesit: {len(text)} karakter.", chart_data)
    except Exception as exc:
        # Show only the exception class, never file contents or local paths.
        return FileResult(name, scope=f"Dosya okunamadı ({type(exc).__name__}). Biçim, bozulma, parola koruması veya Excel çalışma kitabı uyumluluğunu kontrol edin.")


def find_document_matches(text, query, max_results=12, snippet_chars=420):
    """Find bounded, local-only excerpts while retaining page/row/paragraph markers."""
    if not isinstance(text, str) or not isinstance(query, str):
        return []
    terms = search_terms(query)
    if not terms:
        return []
    results, location = [], "Konum belirtilmemiş"
    for line_number, line in enumerate(text.splitlines(), 1):
        stripped = line.strip()
        if not stripped:
            continue
        marker = re.match(r"^\[(Sayfa[^\]]*|Paragraf[^\]]*|Tablo[^\]]*)\]", stripped)
        row_marker = re.match(r"^(Satır \d+:)", stripped)
        if marker:
            location = marker.group(1)
            stripped = stripped[marker.end():].strip()
            if not stripped:
                continue
        if row_marker:
            location = row_marker.group(1)
            stripped = stripped[row_marker.end():].strip()
        normalized = normalize(stripped)
        score = sum(1 for term in terms if term in normalized)
        if score:
            excerpt = stripped
            if len(excerpt) > snippet_chars:
                excerpt = excerpt[:snippet_chars].rstrip() + "…"
            results.append({
                "location": location,
                "line": line_number,
                "excerpt": excerpt,
                "score": score,
            })
    results.sort(key=lambda item: (-item["score"], item["line"]))
    return results[:max(1, min(int(max_results), 30))]


def select_context(results, prompt, limit=5000):
    """Deterministic excerpts, shared fairly; no paid embedding or API calls."""
    valid = [r for r in results if r.text]
    if not valid:
        return ""
    terms = search_terms(prompt)
    chunks = []
    per_file = max(100, (limit - 500) // len(valid))
    for result in valid:
        blocks = document_chunks(result.text)
        scores = [sum(t in normalize(block) for t in terms) for block in blocks]
        # Broad summaries sample the whole indexed document, including its end.
        positions = list(range(len(blocks)))
        if positions and not any(scores):
            anchors = list(dict.fromkeys([0, len(blocks) - 1, len(blocks) // 2, len(blocks) // 4, 3 * len(blocks) // 4]))
            positions = anchors + [i for i in positions if i not in anchors]
        ranked = sorted(positions, key=lambda i: -scores[i])
        selected, remaining = [], per_file
        for index in ranked:
            block = blocks[index]
            if remaining <= 0:
                break
            selected.append((index, block[:remaining]))
            remaining -= len(block)
        chunks.append(f"DOSYA: {result.name[:120]} — {result.status}; ilgi sırasıyla seçilmiş kesitler, tüm belge değil.\n" + "\n[…]\n".join(b for _, b in selected))
    return "\n\n".join(chunks)[:limit]


def normalize(text):
    return text.casefold().translate(str.maketrans("çğıöşü", "cgiosu")).replace("\u0307", "")


def search_terms(prompt):
    stop = {"bir", "icin", "olan", "olarak", "belge", "belgede", "dosya", "dosyada", "bunu", "nedir", "ozetle", "lutfen"}
    return {t for t in re.findall(r"\w{3,}", normalize(prompt)) if t not in stop}


def document_chunks(text):
    """Overlapping chunks retain nearest source marker and exact character range."""
    markers = list(re.finditer(r"\[(?:Sayfa[^\]\n]*|Paragraf[^\]\n]*|Tablo[^\]\n]*)\]|(?:^|\n)Satır \d+:", text))
    blocks, marker_index = [], 0
    for start in range(0, len(text), 780):
        end = min(len(text), start + 900)
        while marker_index + 1 < len(markers) and markers[marker_index + 1].start() <= start:
            marker_index += 1
        location = markers[marker_index].group().strip() if markers and markers[marker_index].start() <= start else ""
        blocks.append(f"[{location}; karakter {start + 1}–{end}]\n" + text[start:end])
    return blocks
