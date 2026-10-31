import io
import json
import os
import unittest
import urllib.error
import zipfile
from unittest.mock import patch

from file_reader import extract_file, select_context, FileResult, MAX_FILE_BYTES, MAX_TEXT, find_document_matches
from model_client import ask_model, read_api_key, build_messages, retry_seconds


class Upload:
    def __init__(self, name, raw):
        self.name, self.raw, self.size = name, raw, len(raw)
    def getvalue(self):
        return self.raw


class CoreTests(unittest.TestCase):
    def test_environment_fallback_when_secrets_missing_key(self):
        with patch.dict(os.environ, {"GROQ_API_KEY": "test-only"}):
            self.assertEqual(read_api_key({}), "test-only")
            self.assertEqual(read_api_key({"GROQ_API_KEY": "secret-test"}), "secret-test")

    def test_missing_key_never_calls_network(self):
        def fail(*a, **k):
            self.fail("Unexpected network call")
        self.assertFalse(ask_model("test", "test", [], opener=fail).ok)

    def test_secrets_failure_falls_back_to_environment(self):
        class BrokenSecrets:
            def get(self, key):
                raise RuntimeError("not configured")
        with patch.dict(os.environ, {"GROQ_API_KEY": "fake"}):
            self.assertEqual(read_api_key(BrokenSecrets()), "fake")

    def test_oversized_prompt_never_calls_network(self):
        def fail(*a, **k):
            self.fail("Unexpected network call")
        self.assertFalse(ask_model("x" * 12001, "test", [], api_key="fake", opener=fail).ok)

    def test_network_error_is_safe(self):
        def fail(*a, **k):
            raise TimeoutError("sensitive-details")
        result = ask_model("test", "test", [], api_key="fake", opener=fail)
        self.assertFalse(result.ok)
        self.assertNotIn("sensitive", result.text)

    def test_rate_limit_honors_retry_and_hides_body(self):
        def limited(*a, **k):
            raise urllib.error.HTTPError("https://example.invalid", 429, "rate", {"Retry-After": "12"}, io.BytesIO(b"sensitive"))
        result = ask_model("test", "test", [], api_key="fake", opener=limited)
        self.assertEqual(result.retry_after, 12)
        self.assertNotIn("sensitive", result.text)

    def test_provider_errors_do_not_leak(self):
        for code in (400, 401, 403, 500):
            def error(*a, **k):
                raise urllib.error.HTTPError("https://example.invalid", code, "error", {}, io.BytesIO(b"sensitive-details"))
            result = ask_model("test", "test", [], api_key="fake", opener=error)
            self.assertFalse(result.ok)
            self.assertNotIn("sensitive", result.text)

    def test_search_payload_and_truncated_response(self):
        def fake(request, **kwargs):
            payload = json.loads(request.data)
            self.assertEqual(payload["tools"], [{"type": "browser_search"}])
            self.assertEqual(payload["tool_choice"], "auto")
            return io.BytesIO(json.dumps({"choices": [{"message": {"content": "[Kaynak](https://example.org/)"}, "finish_reason": "length"}], "usage": {"total_tokens": 10}}).encode())
        result = ask_model("Türkiye nüfus artışı?", "Genel", [], web_enabled=True, api_key="fake", opener=fake)
        self.assertTrue(result.ok)
        self.assertIn("kesildi", result.text)
        self.assertEqual(result.usage["total_tokens"], 10)

    def test_empty_and_invalid_provider_response(self):
        for raw in (b"not json", b'{"choices": []}', b'{"choices":[{"message":{"content":null}}]}'):
            self.assertFalse(ask_model("x", "x", [], api_key="fake", opener=lambda *a, **k: io.BytesIO(raw)).ok)

    def test_whole_recent_messages_and_missing_context_notice(self):
        history = [{"role": "user", "content": "x" * 30000}, {"role": "user", "content": "recent"}]
        result = build_messages("next", "Genel", history)
        self.assertEqual(result[1]["content"], "recent")
        self.assertIn("eski mesajların", result[0]["content"])

    def test_retry_header_variants(self):
        self.assertEqual(retry_seconds("Wed, 01 Jan 2025 00:00:05 GMT", now=1735689600), 5)
        self.assertEqual(retry_seconds("bad"), 60)
        self.assertEqual(retry_seconds("nan"), 60)

    def test_large_file_rejected_before_read(self):
        upload = Upload("big.txt", b"")
        upload.size = MAX_FILE_BYTES + 1
        upload.getvalue = lambda: self.fail("read oversized file")
        self.assertEqual(extract_file(upload).status, "okunamadı")

    def test_text_truncation_visible(self):
        result = extract_file(Upload("sample.txt", b"a" * (MAX_TEXT + 100)))
        self.assertEqual(result.status, "kısmen okundu")
        self.assertEqual(len(result.text), MAX_TEXT)

    def test_csv_semicolon(self):
        result = extract_file(Upload("test.csv", "ad;tutar\nörnek;42".encode()))
        self.assertIn("örnek | 42", result.text)

    def test_corrupt_and_unsupported(self):
        self.assertEqual(extract_file(Upload("test.docx", b"bad")).status, "okunamadı")
        self.assertEqual(extract_file(Upload("test.exe", b"bad")).status, "okunamadı")

    def test_archive_limits(self):
        raw = io.BytesIO()
        with zipfile.ZipFile(raw, "w") as archive:
            for i in range(2001):
                archive.writestr(str(i), "")
        result = extract_file(Upload("test.docx", raw.getvalue()))
        self.assertIn("güvenlik sınırını", result.scope)

    def test_excel_multiple_sheets(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "Bir"
        wb.active.append(["ürün", "adet"])
        wb.active.append(["kalem", 42])
        wb.create_sheet("İki").append(["ikinci", 7])
        raw = io.BytesIO()
        wb.save(raw)
        wb.close()
        result = extract_file(Upload("test.xlsx", raw.getvalue()))
        self.assertIn("kalem", result.text)
        self.assertIn("İki", result.text)
        self.assertEqual(result.status, "kısmen okundu")
        self.assertEqual(result.chart_data["label"], "Bir")
        self.assertEqual(result.chart_data["rows"][0][:2], ["kalem", 42])

    def test_csv_includes_bounded_plot_data(self):
        result = extract_file(Upload("sales.csv", b"urun,tutar\nA,10\nB,20"))
        self.assertEqual(result.chart_data["columns"], ["urun", "tutar"])
        self.assertEqual(result.chart_data["rows"], [["A", "10"], ["B", "20"]])

    def test_image_requires_opt_in_ocr_and_uses_local_reader(self):
        from PIL import Image
        raw = io.BytesIO()
        Image.new("RGB", (20, 20), "white").save(raw, format="PNG")
        upload = Upload("scan.png", raw.getvalue())
        disabled = extract_file(upload, ocr=False)
        self.assertIn("OCR", disabled.scope)
        with patch("ocr_reader.ocr_image", return_value="Makbuz no 123") as reader:
            result = extract_file(upload, ocr=True)
        reader.assert_called_once()
        self.assertIn("Makbuz no 123", result.text)

    def test_pptx_extracts_slide_text_and_table_with_locations(self):
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        box.text = "Çeyrek sonuçları"
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
        table.cell(0, 0).text = "Ürün"
        table.cell(0, 1).text = "Satış"
        table.cell(1, 0).text = "Kalem"
        table.cell(1, 1).text = "24"
        raw = io.BytesIO()
        presentation.save(raw)
        result = extract_file(Upload("sonuc.pptx", raw.getvalue()))
        self.assertIn("[Slayt 1]", result.text)
        self.assertIn("Çeyrek sonuçları", result.text)
        self.assertIn("Kalem | 24", result.text)

    def test_docx_text_and_table(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Örnek rapor")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Tablo verisi"
        raw = io.BytesIO()
        doc.save(raw)
        result = extract_file(Upload("test.docx", raw.getvalue()))
        self.assertIn("Örnek rapor", result.text)
        self.assertIn("Tablo verisi", result.text)

    def test_blank_pdf_detected(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        raw = io.BytesIO()
        writer.write(raw)
        self.assertEqual(extract_file(Upload("blank.pdf", raw.getvalue())).status, "okunamadı")

    def test_document_search_returns_page_and_paragraph_locations(self):
        content = "[Sayfa 2]\n2026 bütçe tablosu ve toplam tutar 1500 TL.\n[Sayfa 3]\nDiğer bilgiler.\n[Paragraf 4]\nSözleşme tarihi 12 Ekim."
        matches = find_document_matches(content, "butce toplam tutar")
        self.assertEqual(len(matches), 1)
        self.assertEqual(matches[0]["location"], "Sayfa 2")
        self.assertIn("1500 TL", matches[0]["excerpt"])

    def test_document_search_supports_excel_rows_and_turkish_normalization(self):
        content = "[Sayfa: Gelir]\n[Sayfa: Gelir, satır 7] A7=İşlem, B7=GÖREV ÜCRETİ, C7=4500"
        matches = find_document_matches(content, "gorev ucreti")
        self.assertEqual(len(matches), 1)
        self.assertEqual(matches[0]["location"], "Sayfa: Gelir, satır 7")
        self.assertIn("4500", matches[0]["excerpt"])

    def test_document_search_is_bounded_and_ignores_empty_query(self):
        content = "\n".join(f"[Sayfa {i}]\nBütçe toplam {i}" for i in range(60))
        self.assertEqual(find_document_matches(content, "   "), [])
        self.assertLessEqual(len(find_document_matches(content, "bütçe", max_results=1000)), 30)

    def test_context_includes_all_files_under_budget(self):
        results = [FileResult(f"file{i}", "a" * 12000, "kısmen okundu") for i in range(5)]
        context = select_context(results, "özetle")
        self.assertLessEqual(len(context), 5000)
        for i in range(5):
            self.assertIn(f"file{i}", context)


if __name__ == "__main__":
    unittest.main()
